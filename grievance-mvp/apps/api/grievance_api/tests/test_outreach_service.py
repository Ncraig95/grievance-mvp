from __future__ import annotations

import base64
import re
import tempfile
from datetime import datetime, timedelta
from io import BytesIO
import logging
from pathlib import Path
import unittest
from zoneinfo import ZoneInfo

from openpyxl import Workbook

from grievance_api.core.config import EmailConfig, OfficerAuthConfig, OutreachConfig
from grievance_api.db.db import Db
from grievance_api.db.migrate import migrate
from grievance_api.services.graph_mail import SentGraphMail
from grievance_api.services.outreach_service import OutreachService


class _FakeMailer:
    def __init__(self):
        self.calls: list[dict[str, object]] = []

    def send_mime_mail(self, **kwargs):  # noqa: ANN003
        self.calls.append(kwargs)
        call_no = len(self.calls)
        return SentGraphMail(
            graph_message_id=f"graph-{call_no}",
            internet_message_id=f"internet-{call_no}",
        )


class _FailingMailer:
    def __init__(self, message: str = "graph permissions denied"):
        self.message = message
        self.calls: list[dict[str, object]] = []

    def send_mime_mail(self, **kwargs):  # noqa: ANN003
        self.calls.append(kwargs)
        raise RuntimeError(self.message)


class OutreachServiceTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.db_path = str(Path(self.tmpdir.name) / "outreach.sqlite3")
        migrate(self.db_path)
        self.db = Db(self.db_path)
        self.email_cfg = EmailConfig(
            enabled=False,
            sender_user_id="",
            templates_dir=str(Path(self.tmpdir.name) / "templates"),
            internal_recipients=(),
            derek_email="grievances@example.org",
            approval_request_url_base="https://grievance.example.org",
            allow_signer_copy_link=False,
            artifact_delivery_mode="sharepoint_link",
            max_attachment_bytes=5_000_000,
            resend_cooldown_seconds=60,
            dry_run=False,
        )
        self.officer_auth_cfg = OfficerAuthConfig(
            enabled=False,
            redirect_uri="https://grievance.example.org/officers/login",
        )

    async def asyncTearDown(self):
        self.tmpdir.cleanup()

    def _service(self, *, enabled: bool, mailer: _FakeMailer | None) -> OutreachService:
        return OutreachService(
            db=self.db,
            logger=logging.getLogger("test.outreach"),
            outreach_cfg=OutreachConfig(
                enabled=enabled,
                sender_user_id="organizing@members.example.org" if enabled else "",
                sender_display_name="CWA Local 3106",
                public_base_url="https://grievance.example.org",
                reply_to_address="replies@example.org",
                reply_to_name="CWA Local 3106",
                timezone="America/New_York",
                min_seconds_between_sends=0.0,
                max_parallel_sends=1,
                max_sends_per_run=50,
            ),
            email_cfg=self.email_cfg,
            officer_auth_cfg=self.officer_auth_cfg,
            mailer=mailer,
        )

    @staticmethod
    def _csv_payload(text: str) -> str:
        return base64.b64encode(text.encode("utf-8")).decode("ascii")

    @staticmethod
    def _xlsx_payload(sheets: dict[str, list[list[object]]]) -> str:
        workbook = Workbook()
        first = True
        for title, rows in sheets.items():
            if first:
                sheet = workbook.active
                sheet.title = title
                first = False
            else:
                sheet = workbook.create_sheet(title)
            for row in rows:
                sheet.append(list(row))
        buf = BytesIO()
        workbook.save(buf)
        workbook.close()
        return base64.b64encode(buf.getvalue()).decode("ascii")

    async def test_seeded_templates_and_preview_use_user_friendly_default_name(self):
        service = self._service(enabled=False, mailer=None)
        await service.ensure_seed_data()

        templates = await service.list_templates()
        stops = await service.list_stops()
        self.assertEqual(len(templates), 2)
        self.assertEqual(len(stops), 18)

        contact = await service.save_contact(
            contact_id=None,
            payload={
                "email": "member@example.org",
                "first_name": "",
                "last_name": "Example",
                "work_location": "Ed Ball Building",
                "active": True,
            },
        )
        notice_template = next(row for row in templates if row["template_type"] == "notice")
        stop = next(row for row in stops if row["location_name"] == "Ed Ball Building" and row["visit_date_local"] == "2026-04-14")

        preview = await service.preview(
            template_id=int(notice_template["id"]),
            stop_id=int(stop["id"]),
            contact_id=int(contact["id"]),
            recipient_email=contact["email"],
        )

        self.assertEqual(preview.subject, "Ed Ball Building Visit on April 14th")
        self.assertIn("Hi everyone,", preview.text_body)
        self.assertIn("https://grievance.example.org/unsubscribe/preview", preview.text_body)
        self.assertEqual(preview.unknown_placeholders, [])

    async def test_run_due_is_idempotent_and_respects_unsubscribe_suppression(self):
        fake_mailer = _FakeMailer()
        service = self._service(enabled=True, mailer=fake_mailer)
        await service.ensure_seed_data()

        await service.save_contact(
            contact_id=None,
            payload={
                "email": "member@example.org",
                "first_name": "Casey",
                "work_location": "Test Hall",
                "active": True,
            },
        )

        zone = ZoneInfo("America/New_York")
        now_local = datetime.now(zone).replace(second=0, microsecond=0)
        due_local = (now_local - timedelta(minutes=10)).strftime("%Y-%m-%dT%H:%M")
        future_local = (now_local + timedelta(days=2)).strftime("%Y-%m-%dT%H:%M")
        visit_date_1 = (now_local.date() + timedelta(days=7)).isoformat()
        visit_date_2 = (now_local.date() + timedelta(days=8)).isoformat()

        await service.save_stop(
            stop_id=None,
            payload={
                "location_name": "Test Hall",
                "visit_date_local": visit_date_1,
                "start_time_local": "10:00",
                "end_time_local": "11:00",
                "timezone": "America/New_York",
                "audience_location": "Test Hall",
                "notice_subject": "Test Hall Visit",
                "reminder_subject": "Reminder Test Hall Visit",
                "notice_send_at_local": due_local,
                "reminder_send_at_local": future_local,
                "status": "active",
            },
        )

        first_run = await service.run_due()
        self.assertEqual(first_run["sent_count"], 1)
        self.assertEqual(first_run["failed_count"], 0)
        self.assertEqual(len(fake_mailer.calls), 1)

        second_run = await service.run_due()
        self.assertEqual(second_run["sent_count"], 0)
        self.assertEqual(second_run["skipped_existing_count"], 1)
        self.assertEqual(len(fake_mailer.calls), 1)

        match = re.search(r"/unsubscribe/([A-Za-z0-9_\\-]+)", str(fake_mailer.calls[0]["text_body"]))
        self.assertIsNotNone(match)
        token = str(match.group(1))
        unsubscribe_result = await service.unsubscribe(token)
        self.assertEqual(unsubscribe_result["email"], "member@example.org")

        await service.save_stop(
            stop_id=None,
            payload={
                "location_name": "Test Hall",
                "visit_date_local": visit_date_2,
                "start_time_local": "10:00",
                "end_time_local": "11:00",
                "timezone": "America/New_York",
                "audience_location": "Test Hall",
                "notice_subject": "Second Test Hall Visit",
                "reminder_subject": "Second Reminder Test Hall Visit",
                "notice_send_at_local": due_local,
                "reminder_send_at_local": future_local,
                "status": "active",
            },
        )

        third_run = await service.run_due()
        self.assertEqual(third_run["sent_count"], 0)
        self.assertEqual(third_run["skipped_suppressed_count"], 2)
        self.assertEqual(len(fake_mailer.calls), 1)

    async def test_send_test_rewrites_links_and_records_analytics(self):
        fake_mailer = _FakeMailer()
        service = self._service(enabled=True, mailer=fake_mailer)
        await service.ensure_seed_data()
        contact = await service.save_contact(
            contact_id=None,
            payload={
                "email": "member@example.org",
                "first_name": "Jordan",
                "work_location": "Analytics Hall",
                "work_group": "Library",
                "active": True,
            },
        )
        template = await service.save_template(
            template_id=None,
            payload={
                "template_key": "analytics_notice",
                "name": "Analytics Notice",
                "template_type": "notice",
                "subject_template": "Analytics Visit",
                "body_template": "Visit https://example.org/join and https://example.org/faq for details.",
                "active": True,
            },
        )
        stop = await service.save_stop(
            stop_id=None,
            payload={
                "location_name": "Analytics Hall",
                "visit_date_local": "2026-06-01",
                "start_time_local": "10:00",
                "end_time_local": "11:00",
                "timezone": "America/New_York",
                "audience_location": "Analytics Hall",
                "notice_subject": "Analytics Visit",
                "reminder_subject": "Analytics Reminder",
                "status": "draft",
            },
        )

        result = await service.send_test(
            template_id=int(template["id"]),
            stop_id=int(stop["id"]),
            contact_id=int(contact["id"]),
            recipient_email=contact["email"],
        )
        self.assertEqual(result.status, "sent")
        self.assertEqual(len(fake_mailer.calls), 1)

        text_body = str(fake_mailer.calls[0]["text_body"])
        html_body = str(fake_mailer.calls[0]["html_body"])
        redirect_match = re.search(r"/r/([A-Za-z0-9_\\-]+)", text_body)
        pixel_match = re.search(r'/o/([^"?]+)\.gif', html_body)
        self.assertIsNotNone(redirect_match)
        self.assertIsNotNone(pixel_match)
        self.assertIn('href="https://grievance.example.org/r/', html_body)

        destination = await service.record_click(str(redirect_match.group(1)), client_ip="203.0.113.10", user_agent="Mozilla/5.0", purpose=None)
        opened = await service.record_open(str(pixel_match.group(1)), client_ip="203.0.113.10", user_agent="Mozilla/5.0", purpose=None)
        dashboard = await service.analytics_dashboard(location="Analytics Hall")

        self.assertEqual(destination, "https://example.org/join")
        self.assertTrue(opened)
        self.assertEqual(dashboard["totals"]["click_count"], 1)
        self.assertEqual(dashboard["totals"]["unique_click_count"], 1)
        self.assertEqual(dashboard["totals"]["estimated_open_count"], 1)
        self.assertEqual(dashboard["totals"]["unique_estimated_open_count"], 1)
        self.assertEqual(dashboard["top_links"][0]["destination_url"], "https://example.org/join")

    async def test_one_off_send_uses_manual_contact_context_without_saved_contact(self):
        fake_mailer = _FakeMailer()
        service = self._service(enabled=True, mailer=fake_mailer)
        await service.ensure_seed_data()
        template = await service.save_template(
            template_id=None,
            payload={
                "template_key": "one_off_notice",
                "name": "One Off Notice",
                "template_type": "notice",
                "subject_template": "Visit for {{ first_name | default('everyone') }}",
                "body_template": "Hi {{ first_name | default('everyone') }}, see us at {{ location }}.",
                "active": True,
            },
        )
        stop = await service.save_stop(
            stop_id=None,
            payload={
                "location_name": "Quick Hall",
                "visit_date_local": "2026-06-03",
                "start_time_local": "12:00",
                "end_time_local": "13:00",
                "timezone": "America/New_York",
                "audience_location": "Quick Hall",
                "notice_subject": "Quick Hall Visit",
                "reminder_subject": "Quick Hall Reminder",
                "status": "draft",
            },
        )

        preview = await service.preview_one_off(
            template_id=int(template["id"]),
            stop_id=int(stop["id"]),
            recipient_email="oneoff@example.org",
            manual_contact={"first_name": "Taylor", "work_group": "Organizing"},
        )
        result = await service.send_one_off(
            template_id=int(template["id"]),
            stop_id=int(stop["id"]),
            recipient_email="oneoff@example.org",
            manual_contact={"first_name": "Taylor", "work_group": "Organizing"},
        )

        self.assertIn("Hi Taylor,", preview.text_body)
        self.assertEqual(result.status, "sent")
        self.assertEqual(len(fake_mailer.calls), 1)
        self.assertIn("Hi Taylor,", str(fake_mailer.calls[0]["text_body"]))

    async def test_quick_message_preview_and_test_send_use_inline_subject_and_body(self):
        fake_mailer = _FakeMailer()
        service = self._service(enabled=True, mailer=fake_mailer)
        await service.ensure_seed_data()
        stop = await service.save_stop(
            stop_id=None,
            payload={
                "location_name": "Message Hall",
                "visit_date_local": "2026-06-04",
                "start_time_local": "08:00",
                "end_time_local": "09:00",
                "timezone": "America/New_York",
                "audience_location": "Message Hall",
                "notice_subject": "Message Hall Visit",
                "reminder_subject": "Message Hall Reminder",
                "status": "draft",
            },
        )

        preview = await service.preview_quick_message(
            stop_id=int(stop["id"]),
            recipient_email="ncraig@cwa3106.com",
            subject_template="Quick test for {{ location }}",
            body_template="Hello from {{ location }} on {{ visit_date }}.",
        )
        result = await service.send_test_quick_message(
            stop_id=int(stop["id"]),
            recipient_email="ncraig@cwa3106.com",
            subject_template="Quick test for {{ location }}",
            body_template="Hello from {{ location }} on {{ visit_date }}.",
        )

        self.assertEqual(preview.subject, "Quick test for Message Hall")
        self.assertIn("Hello from Message Hall", preview.text_body)
        self.assertEqual(result.status, "sent")
        self.assertIn("Quick test for Message Hall", str(fake_mailer.calls[0]["subject"]))

    async def test_send_test_returns_error_text_when_mailer_fails(self):
        failing_mailer = _FailingMailer("mailbox is not permitted to send")
        service = self._service(enabled=True, mailer=failing_mailer)
        await service.ensure_seed_data()
        contact = await service.save_contact(
            contact_id=None,
            payload={"email": "member@example.org", "first_name": "Jordan", "work_location": "Error Hall", "active": True},
        )
        template = await service.save_template(
            template_id=None,
            payload={
                "template_key": "failing_send_notice",
                "name": "Failing Send Notice",
                "template_type": "notice",
                "subject_template": "Failing Send",
                "body_template": "Hello {{ first_name }}.",
                "active": True,
            },
        )
        stop = await service.save_stop(
            stop_id=None,
            payload={
                "location_name": "Error Hall",
                "visit_date_local": "2026-06-07",
                "start_time_local": "10:00",
                "end_time_local": "11:00",
                "timezone": "America/New_York",
                "audience_location": "Error Hall",
                "notice_subject": "Failing Send",
                "reminder_subject": "Failing Send Reminder",
                "status": "draft",
            },
        )

        result = await service.send_test(
            template_id=int(template["id"]),
            stop_id=int(stop["id"]),
            contact_id=int(contact["id"]),
            recipient_email=contact["email"],
        )

        self.assertEqual(result.status, "failed")
        self.assertEqual(result.error_text, "mailbox is not permitted to send")
        send_log = await service.list_send_log()
        self.assertEqual(send_log[0]["status"], "failed")
        self.assertEqual(send_log[0]["error_text"], "mailbox is not permitted to send")

    async def test_inspect_contacts_import_csv_returns_preview_counts_for_combined_status(self):
        service = self._service(enabled=False, mailer=None)
        csv_payload = self._csv_payload(
            "\n".join(
                [
                    "Email Address,First Name,Status",
                    "member1@example.org,Avery,Member - Active - Active",
                    "member2@example.org,Robin,Member - Active - Pending",
                    "member3@example.org,Casey,Member - Inactive - Dropped",
                    ",Taylor,Non Member - Active - Active",
                ]
            )
        )

        result = await service.inspect_contacts_import(
            filename="contacts.csv",
            content_base64=csv_payload,
        )

        self.assertEqual(result["selected_sheet_name"], "CSV")
        self.assertEqual(result["effective_mapping"]["field_mapping"]["email"], "Email Address")
        self.assertEqual(result["effective_mapping"]["status_mapping"]["combined_status_column"], "Status")
        self.assertEqual(result["preview"]["imported_count"], 2)
        self.assertEqual(result["preview"]["updated_count"], 0)
        self.assertEqual(result["preview"]["skipped_reasons"]["missing_email"], 1)
        self.assertEqual(result["preview"]["ignored_reasons"]["unsupported_status"], 1)
        self.assertEqual(result["preview"]["bucket_counts"]["Member - Active - Active"], 1)
        self.assertEqual(result["preview"]["bucket_counts"]["Member - Active - Pending"], 1)

    async def test_inspect_contacts_import_xlsx_supports_sheet_selection(self):
        service = self._service(enabled=False, mailer=None)
        payload = self._xlsx_payload(
            {
                "Ignore Me": [
                    ["Email", "Status"],
                    ["ignored@example.org", "Member - Inactive - Dropped"],
                ],
                "Target Sheet": [
                    ["Email", "Status"],
                    ["member@example.org", "Non Member - Active - Active"],
                ],
            }
        )

        result = await service.inspect_contacts_import(
            filename="contacts.xlsx",
            content_base64=payload,
            sheet_name="Target Sheet",
        )

        self.assertEqual([sheet["name"] for sheet in result["sheets"]], ["Ignore Me", "Target Sheet"])
        self.assertEqual(result["selected_sheet_name"], "Target Sheet")
        self.assertEqual(result["preview"]["imported_count"], 1)
        self.assertEqual(result["preview"]["bucket_counts"]["Non Member - Active - Active"], 1)

    async def test_import_contacts_remembers_confirmed_mapping_profile(self):
        service = self._service(enabled=False, mailer=None)
        csv_payload = self._csv_payload(
            "\n".join(
                [
                    "Email Address,Worker First,Worker Last,Rollup Status",
                    "member@example.org,Jamie,Lee,Member - Active - Active",
                ]
            )
        )
        mapping = {
            "field_mapping": {
                "email": "Email Address",
                "first_name": "Worker First",
                "last_name": "Worker Last",
            },
            "status_mapping": {
                "mode": "combined",
                "combined_status_column": "Rollup Status",
            },
        }

        first_import = await service.import_contacts(
            filename="mapping.csv",
            content_base64=csv_payload,
            mapping=mapping,
        )
        second_inspect = await service.inspect_contacts_import(
            filename="mapping.csv",
            content_base64=csv_payload,
        )

        self.assertTrue(first_import["saved_mapping"])
        self.assertIsNotNone(second_inspect["remembered_mapping"])
        self.assertEqual(second_inspect["remembered_mapping"]["field_mapping"]["first_name"], "Worker First")
        self.assertEqual(second_inspect["effective_mapping"]["field_mapping"]["last_name"], "Worker Last")
        self.assertEqual(second_inspect["effective_mapping"]["status_mapping"]["combined_status_column"], "Rollup Status")

    async def test_import_contacts_combined_status_only_imports_allowed_buckets(self):
        service = self._service(enabled=False, mailer=None)
        csv_payload = self._csv_payload(
            "\n".join(
                [
                    "Email,Status,Work Location",
                    "member-active@example.org,Member - Active - Active,North Hall",
                    "member-pending@example.org,Member - Active - Pending,North Hall",
                    "member-inactive@example.org,Member - Inactive - Dropped,North Hall",
                    "nonmember-active@example.org,Non Member - Active - Active,North Hall",
                    "nonmember-nonfr@example.org,Non Member - Active - Non fr Mem,North Hall",
                ]
            )
        )

        result = await service.import_contacts(
            filename="combined.csv",
            content_base64=csv_payload,
        )
        contacts = await service.list_contacts()

        self.assertEqual(result["imported_count"], 4)
        self.assertEqual(result["ignored_count"], 1)
        self.assertEqual(result["ignored_reasons"]["unsupported_status"], 1)
        self.assertEqual(len(contacts), 4)
        self.assertEqual(
            sorted(contact["status_bucket"] for contact in contacts),
            sorted(
                [
                    "Member - Active - Active",
                    "Member - Active - Pending",
                    "Non Member - Active - Active",
                    "Non Member - Active - Non fr Mem",
                ]
            ),
        )

    async def test_import_contacts_split_status_only_imports_allowed_buckets(self):
        service = self._service(enabled=False, mailer=None)
        csv_payload = self._csv_payload(
            "\n".join(
                [
                    "Email,Membership Type,Employment Status,Status Detail",
                    "member-active@example.org,Member,Active,Active",
                    "member-pending@example.org,Member,Active,Pending",
                    "member-inactive@example.org,Member,Inactive,Dropped",
                    "nonmember-active@example.org,Non Member,Active,Active",
                    "nonmember-nonfr@example.org,Non Member,Active,Non fr Mem",
                ]
            )
        )
        mapping = {
            "field_mapping": {
                "email": "Email",
            },
            "status_mapping": {
                "mode": "split",
                "membership_type_column": "Membership Type",
                "employment_status_column": "Employment Status",
                "status_detail_column": "Status Detail",
            }
        }

        result = await service.import_contacts(
            filename="split.csv",
            content_base64=csv_payload,
            mapping=mapping,
        )
        contacts = await service.list_contacts()

        self.assertEqual(result["imported_count"], 4)
        self.assertEqual(result["ignored_count"], 1)
        self.assertEqual(result["ignored_reasons"]["unsupported_status"], 1)
        self.assertEqual(len(contacts), 4)
        saved = {contact["email"]: contact for contact in contacts}
        self.assertEqual(saved["member-pending@example.org"]["status_bucket"], "Member - Active - Pending")
        self.assertEqual(saved["nonmember-nonfr@example.org"]["status_bucket"], "Non Member - Active - Non fr Mem")

    async def test_contacts_for_stop_filters_by_audience_status_bucket(self):
        service = self._service(enabled=False, mailer=None)
        await service.save_contact(
            contact_id=None,
            payload={
                "email": "member-active@example.org",
                "first_name": "Avery",
                "work_location": "Target Hall",
                "active": True,
                "membership_type": "Member",
                "employment_status": "Active",
                "status_detail": "Active",
                "status_bucket": "Member - Active - Active",
                "status_source_text": "Member - Active - Active",
            },
        )
        await service.save_contact(
            contact_id=None,
            payload={
                "email": "member-pending@example.org",
                "first_name": "Robin",
                "work_location": "Target Hall",
                "active": True,
                "membership_type": "Member",
                "employment_status": "Active",
                "status_detail": "Pending",
                "status_bucket": "Member - Active - Pending",
                "status_source_text": "Member - Active - Pending",
            },
        )
        stop = await service.save_stop(
            stop_id=None,
            payload={
                "location_name": "Target Hall",
                "visit_date_local": "2026-07-01",
                "start_time_local": "09:00",
                "end_time_local": "10:00",
                "timezone": "America/New_York",
                "audience_location": "Target Hall",
                "audience_status_bucket": "Member - Active - Pending",
                "notice_subject": "Target Hall Visit",
                "reminder_subject": "Target Hall Reminder",
                "status": "active",
            },
        )

        contacts = await service._contacts_for_stop(stop)  # noqa: SLF001

        self.assertEqual([row["email"] for row in contacts], ["member-pending@example.org"])

    async def test_prefetch_click_is_flagged_and_excluded_from_topline_metrics(self):
        fake_mailer = _FakeMailer()
        service = self._service(enabled=True, mailer=fake_mailer)
        await service.ensure_seed_data()
        contact = await service.save_contact(
            contact_id=None,
            payload={"email": "prefetch@example.org", "first_name": "Scout", "work_location": "Prefetch Hall", "active": True},
        )
        template = await service.save_template(
            template_id=None,
            payload={
                "template_key": "prefetch_notice",
                "name": "Prefetch Notice",
                "template_type": "notice",
                "subject_template": "Prefetch Visit",
                "body_template": "Open https://example.org/prefetch now.",
                "active": True,
            },
        )
        stop = await service.save_stop(
            stop_id=None,
            payload={
                "location_name": "Prefetch Hall",
                "visit_date_local": "2026-06-02",
                "start_time_local": "09:00",
                "end_time_local": "10:00",
                "timezone": "America/New_York",
                "audience_location": "Prefetch Hall",
                "notice_subject": "Prefetch Visit",
                "reminder_subject": "Prefetch Reminder",
                "status": "draft",
            },
        )

        await service.send_test(
            template_id=int(template["id"]),
            stop_id=int(stop["id"]),
            contact_id=int(contact["id"]),
            recipient_email=contact["email"],
        )
        text_body = str(fake_mailer.calls[0]["text_body"])
        redirect_match = re.search(r"/r/([A-Za-z0-9_\\-]+)", text_body)
        self.assertIsNotNone(redirect_match)

        await service.record_click(
            str(redirect_match.group(1)),
            client_ip="203.0.113.99",
            user_agent="Mozilla/5.0 urlscan security crawler",
            purpose="prefetch",
        )
        dashboard = await service.analytics_dashboard(location="Prefetch Hall")

        self.assertEqual(dashboard["totals"]["click_count"], 0)
        self.assertEqual(dashboard["totals"]["unique_click_count"], 0)
        self.assertTrue(dashboard["recent_activity"][0]["suspected_automation"])


if __name__ == "__main__":
    unittest.main()
